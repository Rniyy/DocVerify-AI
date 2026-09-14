from docx import Document


def extract_docx(path: str) -> dict:
    """Pulls paragraph text and table cell text out of a .docx file.

    Tables are flattened to pipe-separated rows so a "Label | Value" style
    table cell still lines up as a single line the same label/value line
    matcher used for PDFs (Node's mapTextToFields) can parse.
    """
    document = Document(path)
    parts = [p.text for p in document.paragraphs if p.text.strip()]

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))

    return {"type": "docx", "text": "\n".join(parts)}
